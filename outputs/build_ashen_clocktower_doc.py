from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"D:\UNITY2\LearnMetroidvania\Metroidvania\outputs\灰烬钟塔项目技术点_腾讯游戏客户端实习面试.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 30, 42)
GRAY = RGBColor(90, 90, 90)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_east_asia_font(run, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def style_run(run, size=11, bold=None, italic=None, color=None):
    set_east_asia_font(run)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph(p, before=0, after=6, line=1.1, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_p(doc, text="", size=11, bold=False, italic=False, color=INK, after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style)
    set_paragraph(p, before=before, after=after, line=1.1, align=align)
    if text:
        run = p.add_run(text)
        style_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_h(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph(p, after=4, line=1.167)
        run = p.add_run(item)
        style_run(run, size=11, color=INK)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph(p, after=4, line=1.167)
        run = p.add_run(item)
        style_run(run, size=11, color=INK)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    set_paragraph(p, after=2, line=1.1)
    style_run(p.add_run(title + "："), size=11, bold=True, color=DARK_BLUE)
    style_run(p.add_run(body), size=11, color=INK)
    add_p(doc, "", after=4)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.35, 5.15])
    for i, (k, v) in enumerate(rows):
        set_cell_shading(table.cell(i, 0), LIGHT_GRAY)
        for cell in table.row_cells(i):
            p = cell.paragraphs[0]
            set_paragraph(p, after=0, line=1.1)
        style_run(table.cell(i, 0).paragraphs[0].add_run(k), size=10.5, bold=True, color=DARK_BLUE)
        style_run(table.cell(i, 1).paragraphs[0].add_run(v), size=10.5, color=INK)
    return table


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for i, header in enumerate(headers):
        set_cell_shading(table.cell(0, i), LIGHT_GRAY)
        p = table.cell(0, i).paragraphs[0]
        set_paragraph(p, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_run(p.add_run(header), size=10, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_paragraph(p, after=0, line=1.1)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(value), size=9.5, color=INK)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, attr, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    set_paragraph(header_p, after=0, line=1.0)
    style_run(header_p.add_run("灰烬钟塔项目技术复盘"), size=9.5, color=GRAY)
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    set_paragraph(footer_p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(footer_p.add_run("腾讯游戏客户端实习面试准备"), size=9.5, color=GRAY)


def build():
    doc = Document()
    configure_document(doc)

    add_p(doc, "项目技术复盘", size=12, bold=True, color=GRAY, after=2)
    add_p(doc, "灰烬钟塔 Ashen Clocktower", size=25, bold=True, color=RGBColor(0, 0, 0), after=4)
    add_p(doc, "面向腾讯游戏客户端实习面试的技术点梳理", size=14, color=GRAY, after=14)
    add_kv_table(doc, [
        ("项目类型", "Unity 2022.3 / 2D Metroidvania 横版动作探索原型"),
        ("核心目标", "在已有角色、战斗、技能、UI、存档框架上，扩展一套围绕“灰烬相位”的关卡与叙事玩法。"),
        ("推荐讲法", "先讲运行时系统，再讲编辑器工具链，最后讲工程取舍和可扩展性。"),
        ("面试定位", "突出 Unity 客户端能力：状态机、Tilemap/Collider、UGUI/TMP、Shader、EditorWindow、存档与模块解耦。"),
    ])
    add_callout(doc, "一句话项目介绍", "灰烬钟塔是一个 2D 类银河恶魔城关卡扩展：玩家消耗灰烬值进入“灰烬相位”，让过去的桥梁、平台、机关和叙事线索短暂显现，并通过能力门、传送门、Boss 前流程形成完整探索闭环。")

    add_h(doc, "1. 项目整体技术架构", 1)
    add_p(doc, "灰烬钟塔并不是孤立关卡，而是接入原项目多套系统后的玩法扩展。原项目已有 Entity、Player/Enemy FSM、Skill、Stats、Buff、Inventory、Audio、UI、Save&Load 等模块；灰烬钟塔主要新增 Ash、Narrative、MapGeneration、Editor Tool 四条线。")
    add_matrix(doc, ["模块", "关键职责", "面试可讲点"], [
        ("角色/敌人 FSM", "Player/Enemy 通过状态类封装 Idle、Move、Jump、Dash、Attack、Dead 等行为。", "用状态机降低 Update 分支复杂度，动画参数和逻辑状态一一对应。"),
        ("灰烬相位", "AshSystem 管资源，PhaseShiftController 管切相，AshenObject 管显隐与碰撞。", "资源、表现、关卡实体分层，避免把玩法逻辑写死在地图对象里。"),
        ("地图生成工具", "EditorWindow 生成关卡占位、房间蓝图、Tilemap 预览和完整 Tilemap 测试地图。", "用编辑器工具把重复搭建流程自动化，减少手工摆放错误。"),
        ("叙事/UI", "NarrativeTrigger、DialogueManager、DialogueUI、FullscreenTextSequence 负责触发、打字机播放和全屏文本。", "交互触发、自动触发和相位限定触发共用一套播放链路。"),
        ("存档与能力", "PlayerManager 与 GameData 保存能力解锁，OldDayBlessing 同步到 AshSystem。", "能力状态持久化，重开游戏后灰烬机制保持一致。"),
    ], [1.25, 2.35, 2.9])

    add_h(doc, "2. 灰烬相位系统", 1)
    add_p(doc, "灰烬相位是本项目最适合重点展开的技术点。它把“资源消耗”和“世界状态切换”结合起来：玩家按 R 进入相位，灰烬值持续下降；相位开启时，隐藏的 Tilemap、平台、Collider、剧情触发和视觉滤镜同时生效；灰烬耗尽后自动退出。")
    add_h(doc, "2.1 运行时职责拆分", 2)
    add_bullets(doc, [
        "AshSystem：维护 currentAsh/maxAsh，提供 AddAsh、ConsumeAsh、HasEnoughAsh、IsEmpty，并通过 OnAshChanged 事件驱动 UI。",
        "PhaseShiftController：监听 R 键，判断进入门槛，持续按 Time.deltaTime 消耗灰烬，并在耗尽时退出。",
        "AshenWorldManager：统一查找场景中的 AshenObject，用 SetAshenWorldActive 批量切换灰烬世界显隐。",
        "AshenObject：缓存子节点 Renderer 与 Collider2D，切相时只开关渲染和碰撞，不 SetActive，避免管理器找不到隐藏对象。",
        "AshenPhaseShaderController：运行时创建 ScreenSpaceOverlay、材质和粒子，做相位滤镜淡入淡出。"
    ])
    add_h(doc, "2.2 技术难点与解决", 2)
    add_matrix(doc, ["问题", "方案", "收益"], [
        ("隐藏平台既要不可见，也要不可碰撞", "AshenObject 同时控制 Renderer.enabled 和 Collider2D.enabled。", "表现和物理一致，避免玩家踩到不可见平台。"),
        ("隐藏对象不能因 SetActive(false) 被查找遗漏", "对象本体保持激活，仅关闭组件；FindObjectsOfType<AshenObject>(true) 做兜底刷新。", "关卡设计师可直接在层级里管理灰烬对象，运行时也稳定。"),
        ("相位视觉要覆盖全屏且不影响输入", "ScreenSpaceOverlay + Image + 自定义 Shader，raycastTarget=false。", "视觉层独立，不卡住 UI/角色输入。"),
        ("资源消耗要和帧率无关", "ConsumeAsh(ashDrainPerSecond * Time.deltaTime)。", "不同机器上相位持续时间一致。"),
        ("后期能力可能改变资源规则", "OldDayBlessing 通过 AshSystem.SetOldDayBlessing 跳过消耗判断。", "祝福能力无需改切相代码，扩展点清晰。"),
    ], [1.8, 2.6, 2.1])
    add_callout(doc, "面试表达", "我把灰烬系统拆成资源、相位控制、世界对象、视觉表现四层。这样单个模块的职责比较窄，后续想加 UI、音效、能力免耗、特殊灰烬机关，都不需要把 PhaseShiftController 改成一个大脚本。")

    add_h(doc, "3. Shader 与视觉反馈", 1)
    add_p(doc, "灰烬相位的表现不是简单调色，而是运行时创建覆盖层材质，通过 Shader 参数控制滤镜颜色、强度、暗度和灰烬噪声。控制器使用协程在 0.25 秒内插值强度，并配合粒子系统表现漂浮灰尘。")
    add_bullets(doc, [
        "Shader 使用 Queue=Transparent、ZTest Always、Blend SrcAlpha OneMinusSrcAlpha，保证滤镜覆盖在画面上方。",
        "片元阶段用 Hash21 生成基于 UV 网格的灰烬噪声，叠加 _AshSpeed 形成轻微流动感。",
        "控制器缓存 Shader.PropertyToID，减少每帧字符串查找。",
        "粒子系统可手动绑定，也可缺失时运行时创建并挂在主相机下，降低场景装配成本。"
    ])

    add_h(doc, "4. 编辑器地图生成工具链", 1)
    add_p(doc, "灰烬钟塔有多套 Editor 工具，重点价值是把“关卡流程、占位、Tilemap 预览、最终测试地图”分阶段生成。对游戏客户端实习面试来说，这能体现 Unity Editor 扩展、关卡生产效率和工程化意识。")
    add_matrix(doc, ["工具", "做什么", "技术点"], [
        ("AshenMapGeneratorWindow", "一键生成完整地图骨架：区域、普通平台、灰烬平台、能力门、敌人占位、存档点、传送门、叙事触发。", "EditorWindow、Undo、Scene dirty、反射查找组件、颜色占位和 Gizmo 标记。"),
        ("AshenRoomBlueprintWindow", "按房间生成 27 个左右的房间蓝图，并导出房间说明文本。", "数据驱动房间定义，RoomType 决定尺寸、平台、敌人、能力点和教学重点。"),
        ("AshenRoomTilemapPreviewWindow", "生成 Ground/Ashen 双 Tilemap 预览，并可实例化传送门、篝火 Prefab。", "Grid/Tilemap/TilemapCollider2D/CompositeCollider2D 自动装配。"),
        ("AshenClocktowerTestMapGenerator", "生成较完整的 Tilemap 测试地图，拆分 Background、Details、Ground、Walls、Hazards、Ashen 层。", "Tile 资源查找、坐标日志导出、多层 Tilemap 碰撞分离。"),
    ], [1.65, 2.65, 2.2])
    add_h(doc, "4.1 工具链亮点", 2)
    add_bullets(doc, [
        "所有生成物都挂在固定 Root 下，删除和重生成只影响工具产物，不破坏场景中已有手工内容。",
        "使用 Undo.RegisterCreatedObjectUndo 和 Undo.DestroyObjectImmediate，符合 Unity 编辑器操作习惯。",
        "把能力门、敌人、机关先做成 GeneratedMapMarker/GeneratedAbilityGate 占位，避免早期过度绑定正式玩法脚本。",
        "Tilemap 预览自动添加 Rigidbody2D、TilemapCollider2D 和 CompositeCollider2D，减少手动碰撞配置遗漏。",
        "导出坐标清单，便于从灰盒关卡过渡到正式美术/Prefab 替换。"
    ])
    add_callout(doc, "面试表达", "我不是直接堆一个最终地图，而是做了从设计蓝图到 Tilemap 预览再到完整测试地图的工具链。这样设计调整时成本更低，也方便把地图坐标、能力门和叙事点位交给后续正式资源替换。")

    add_h(doc, "5. 叙事触发与 UI 播放", 1)
    add_p(doc, "灰烬钟塔的叙事系统服务于探索节奏：进入区域自动触发背景文本，靠近石碑等对象时按 E 交互触发，部分文本还要求玩家处于灰烬相位。")
    add_matrix(doc, ["组件", "职责", "设计取舍"], [
        ("NarrativeTrigger", "检测玩家进入/离开、是否需要交互键、是否要求灰烬相位。", "把触发条件放在触发器脚本里，DialogueManager 只负责播放。"),
        ("DialogueManager", "单例管理播放状态，防止多个文本互相覆盖。", "IsPlaying 作为全局互斥，简单但有效。"),
        ("DialogueUI", "TMP 打字机显示，支持空格跳过当前句/下一句，自动触发文本延迟后按 E 继续。", "自动触发和交互触发用不同继续键与提示，符合操作语义。"),
        ("FullscreenTextSequence", "开场/结局全屏文字，播放时可暂停 Time.timeScale。", "剧情演出不受角色移动和战斗干扰。"),
    ], [1.45, 2.45, 2.6])

    add_h(doc, "6. 能力解锁、存档与灰烬规则", 1)
    add_p(doc, "项目原本已有能力解锁与存档框架。灰烬钟塔新增“旧日祝福”能力后，没有单独再造一套存档，而是接入 PlayerManager 和 GameData：获得祝福时设置 ability_HasOldDayBlessing，并同步 AshSystem 的 hasOldDayBlessing；存档/读档时也恢复该状态。")
    add_bullets(doc, [
        "PlayerManager.SaveData/LoadData 负责保存能力布尔值，GameData 增加 hasOldDayBlessing 字段。",
        "OldDayBlessing 触发器处理玩家进入范围、E 键交互、一次性授予、音效和提示隐藏。",
        "AshRewardStatue 支持攻击获得灰烬，并限制单轮最大可奖励次数；保存游戏后通过 SavesManager.OnGameSaved 重置命中次数。",
        "AshHealController 复用 AshSystem 的资源判断与消耗，满血时不消耗资源，避免交互浪费。"
    ])

    add_h(doc, "7. 可重点讲的工程取舍", 1)
    add_numbers(doc, [
        "运行时系统按职责拆分，而不是把资源、视觉、Collider、UI 全塞进一个玩家脚本。",
        "Editor 工具优先生成可编辑占位对象，避免把早期关卡设计锁死在最终 Prefab 上。",
        "灰烬平台通过组件启停而不是 GameObject 启停，解决隐藏对象查找和运行时切换的一致性问题。",
        "叙事系统允许自动触发、交互触发、相位限定触发共存，但播放层保持统一。",
        "能力和祝福复用存档框架，说明新增玩法时会优先接入已有架构，而不是重复造轮子。"
    ])

    add_h(doc, "8. 面试高频追问准备", 1)
    add_matrix(doc, ["追问", "建议回答"], [
        ("如果灰烬对象很多，FindObjectsOfType 会不会有性能问题？", "当前只在 Start/Refresh 时查找，不在每帧查找。规模变大后可以改成 AshenObject OnEnable 注册到 Manager，或按区域分组管理。"),
        ("为什么不用 SetActive 控制灰烬平台？", "SetActive 会让隐藏对象从常规查找和生命周期里消失，且子对象脚本状态可能被打断。这里仅开关 Renderer/Collider，更适合短时间相位切换。"),
        ("如何避免玩家在灰烬平台消失时卡住？", "现在通过灰烬值耗尽退出，需要进一步做离开前检测：例如退出时检查玩家脚下是否只依赖灰烬平台，给短暂缓冲或强制传送到最近安全点。"),
        ("工具生成地图和正式地图如何衔接？", "生成物使用固定命名、坐标日志、Marker 和 AbilityGate 记录设计意图，后续可以替换为正式 Prefab 或用脚本批量转化。"),
        ("哪些地方还能优化？", "灰烬对象注册制、对象池化粒子/奖励反馈、ScriptableObject 化房间数据、能力门正式化、存档版本兼容、自动化 PlayMode 测试。"),
    ], [2.0, 4.5])

    add_h(doc, "9. 可扩展方向", 1)
    add_bullets(doc, [
        "把房间定义从硬编码迁移到 ScriptableObject，允许策划在 Inspector 中编辑房间结构和教学目标。",
        "为 AshenObject 增加区域 ID，配合摄像机/房间触发只刷新当前区域，减少大地图切相成本。",
        "能力门从设计占位升级为正式组件，读取 PlayerManager 能力状态并控制门、UI 提示和音效。",
        "把灰烬消耗、回血消耗、奖励数值抽成配置表，方便平衡性迭代。",
        "补充 PlayMode 测试：切相显隐、灰烬耗尽退出、祝福免耗、叙事互斥播放、存档恢复能力。"
    ])

    add_h(doc, "10. 源码对应关系", 1)
    add_matrix(doc, ["技术点", "主要文件"], [
        ("灰烬资源与相位", "Assets/Scripts/Ash/AshSystem.cs, PhaseShiftController.cs, AshenWorldManager.cs, AshenObject.cs"),
        ("视觉滤镜与粒子", "Assets/Scripts/Ash/AshenPhaseShaderController.cs, Assets/Shaders/Ash/AshenPhaseFilter.shader"),
        ("编辑器地图工具", "Assets/Editor/AshenClocktowerMapTool/*.cs, Assets/Editor/AshenClocktowerTestMapGenerator.cs"),
        ("叙事播放", "Assets/Scripts/Narrative/NarrativeTrigger.cs, DialogueManager.cs, DialogueUI.cs, FullscreenTextSequence.cs"),
        ("能力/存档接入", "Assets/Script/Manager/PlayerManager.cs, Assets/Script/Save&Load/GameData.cs, Assets/Script/Interact/OldDayBlessing.cs"),
    ], [1.8, 4.7])

    add_callout(doc, "最后建议", "面试时不要从脚本清单开始讲。先讲“灰烬相位解决了什么体验问题”，再展开运行时系统和编辑器工具链，最后主动讲一个你已经意识到的优化点，会更像真实客户端开发。")

    doc.save(OUT)


if __name__ == "__main__":
    build()
